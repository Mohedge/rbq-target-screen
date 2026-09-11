"""
Acquisition target screen: RBQ active licences x REQ business register.

Deterministic pipeline. Invents nothing. Every filtering step reports its row
count; an empty result stops and names the filter that emptied it.
"""

import io
import unicodedata
import zipfile
from dataclasses import dataclass, field, asdict

import pandas as pd

# ---------------------------------------------------------------------------
# Parameters
# ---------------------------------------------------------------------------

@dataclass
class Params:
    SUBCATEGORY: str = "16"                 # 16 = electrical; change for another trade
    MUNICIPALITIES: list = field(default_factory=lambda: [
        # Placeholder corridor (Laurentides, Route 117 / A-15). Edit freely.
        "Saint-Sauveur", "Piedmont", "Sainte-Adèle", "Sainte-Anne-des-Lacs",
        "Prévost", "Saint-Hippolyte", "Morin-Heights", "Val-Morin", "Val-David",
        "Sainte-Agathe-des-Monts", "Saint-Jérôme", "Saint-Colomban",
        "Sainte-Sophie", "Mirabel", "Saint-Adolphe-d'Howard",
        "Sainte-Marguerite-du-Lac-Masson",
    ])
    CURRENT_YEAR: int = 2026
    AGE_THRESHOLD: int = 15
    STRONG_AGE_THRESHOLD: int = 25
    HIGH_BOND: int = 40000
    SUBCAT_THRESHOLD: int = 10
    OUTPUT_ROWS: int = 40
    USE_REQ_JOIN: bool = True


class EmptyResult(Exception):
    """Raised when a filter empties the dataset. Message names the filter."""


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s)
                   if not unicodedata.combining(c))


def normalize_headers(df: pd.DataFrame) -> pd.DataFrame:
    """Accent-insensitive headers so 'Sous-catégories' == 'Sous-categories'."""
    df.columns = [strip_accents(c).strip().lstrip("\ufeff") for c in df.columns]
    return df


def norm_key(s: str) -> str:
    """Accent- and case-insensitive comparison key for municipality names."""
    return strip_accents(str(s)).strip().casefold()


def open_csv_from_zip(zip_bytes: bytes, member_hint: str) -> io.BytesIO:
    """Return a file-like object for the first zip member whose name contains
    member_hint (case-insensitive). Reads only that member."""
    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    names = [n for n in zf.namelist() if member_hint.lower() in n.lower()]
    if not names:
        raise FileNotFoundError(f"No member matching '{member_hint}' in zip: {zf.namelist()}")
    return io.BytesIO(zf.read(names[0]))


# ---------------------------------------------------------------------------
# Steps
# ---------------------------------------------------------------------------

def step1_read_rbq(source, p: Params, log, chunksize=200_000) -> pd.DataFrame:
    """source: path, file-like, or bytes of the CSV. Every column as text.

    Reads in chunks and applies the subcategory filter (step 2a) per chunk, so
    the full 900k-row file never sits in memory at once. The raw row count is
    still reported as STEP 1; the filtered count is reported as STEP 2a."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    total = 0
    parts = []
    reader = pd.read_csv(source, dtype=str, keep_default_na=False,
                         encoding="utf-8", chunksize=chunksize)
    for chunk in reader:
        chunk = normalize_headers(chunk)
        total += len(chunk)
        parts.append(chunk[chunk["Sous-categories"].str.strip() == p.SUBCATEGORY])
    df = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
    del parts
    log("STEP 1 - READ", total)
    log(f"STEP 2a - Sous-categories == {p.SUBCATEGORY}", len(df))
    if df.empty:
        raise EmptyResult(f"Sous-categories == {p.SUBCATEGORY}")
    return df


def step2_filter(df: pd.DataFrame, p: Params, log) -> pd.DataFrame:

    df = df[df["Type de licence"].str.strip() == "Entrepreneur"]
    log("STEP 2b - Type de licence == Entrepreneur", len(df))
    if df.empty:
        raise EmptyResult("Type de licence == Entrepreneur")

    df = df[df["Statut juridique"].str.strip().isin(["Personne morale", "Corporation"])]
    log("STEP 2c - Statut juridique in {Personne morale, Corporation}", len(df))
    if df.empty:
        raise EmptyResult("Statut juridique in {Personne morale, Corporation}")

    df = df.drop_duplicates(subset="Numero de licence")
    log("STEP 2d - Deduplicate on Numero de licence", len(df))
    return df


def step3_territory(df: pd.DataFrame, p: Params, log) -> pd.DataFrame:
    keys = {norm_key(m) for m in p.MUNICIPALITIES}
    df = df[df["Municipalite"].map(norm_key).isin(keys)]
    log("STEP 3 - Municipalite in MUNICIPALITIES", len(df))
    if df.empty:
        raise EmptyResult("Municipalite in MUNICIPALITIES")
    return df


def step4_join_req(df: pd.DataFrame, nom_source, log, chunksize=500_000) -> pd.DataFrame:
    """Read Nom.csv in chunks, keeping only NEQs in df; earliest DAT_INIT_NOM_ASSUJ
    per NEQ becomes the registration year. Memory stays flat regardless of file size."""
    if isinstance(nom_source, (bytes, bytearray)):
        nom_source = io.BytesIO(nom_source)
    wanted = set(df["NEQ"].str.strip())
    wanted.discard("")
    parts = []
    reader = pd.read_csv(
        nom_source, dtype=str, keep_default_na=False, encoding="utf-8-sig",
        usecols=["NEQ", "NOM_ASSUJ", "DAT_INIT_NOM_ASSUJ"], chunksize=chunksize,
    )
    for chunk in reader:
        parts.append(chunk[chunk["NEQ"].isin(wanted)])
    nom = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame(
        columns=["NEQ", "NOM_ASSUJ", "DAT_INIT_NOM_ASSUJ"])
    del parts

    nom["dt"] = pd.to_datetime(nom["DAT_INIT_NOM_ASSUJ"], errors="coerce")
    earliest = (nom.dropna(subset=["dt"])
                   .sort_values("dt")
                   .drop_duplicates(subset="NEQ", keep="first")
                   [["NEQ", "dt"]])
    earliest["Registration year"] = earliest["dt"].dt.year.astype("Int64")
    earliest = earliest.drop(columns="dt")

    df = df.merge(earliest, on="NEQ", how="left")
    matched = df["Registration year"].notna().sum()
    rate = 100.0 * matched / len(df) if len(df) else 0.0
    log(f"STEP 4 - REQ match rate: {matched}/{len(df)} = {rate:.1f}%", len(df))
    return df


def step5_calculate(df: pd.DataFrame, p: Params, log) -> pd.DataFrame:
    df = df.copy()
    if "Registration year" not in df.columns:
        df["Registration year"] = pd.Series([pd.NA] * len(df), dtype="Int64")
    df["Age"] = (p.CURRENT_YEAR - df["Registration year"]).astype("Int64")
    df["Bond"] = pd.to_numeric(df["Montant de la caution"], errors="coerce")
    df["Subcategory count"] = pd.to_numeric(
        df["Nombre de sous-categorie autorisees"], errors="coerce")
    log("STEP 5 - CALCULATE", len(df))
    return df


def step6_score(df: pd.DataFrame, p: Params, log) -> pd.DataFrame:
    df = df.copy()
    age = df["Age"].astype("float")
    score = pd.Series(0, index=df.index, dtype="int64")
    score += (age >= p.STRONG_AGE_THRESHOLD).fillna(False).astype(int) * 3
    score += ((age >= p.AGE_THRESHOLD) & (age < p.STRONG_AGE_THRESHOLD)).fillna(False).astype(int) * 2
    score += (df["Bond"] == p.HIGH_BOND).fillna(False).astype(int) * 2
    score += (df["Subcategory count"] >= p.SUBCAT_THRESHOLD).fillna(False).astype(int) * 1
    df["Score"] = score
    log("STEP 6 - SCORE", len(df))
    return df


def step7_rank(df: pd.DataFrame, p: Params, log) -> pd.DataFrame:
    df = df.sort_values(
        ["Score", "Registration year"], ascending=[False, True], na_position="last")
    shortlist = df.head(p.OUTPUT_ROWS)
    log(f"STEP 7 - RANK, shortlist top {p.OUTPUT_ROWS}", len(shortlist))
    return shortlist


OUTPUT_MAP = [
    ("Company", "Nom de l'intervenant"),
    ("NEQ", "NEQ"),
    ("Municipality", "Municipalite"),
    ("Region", "Region administrative"),
    ("Licence number", "Numero de licence"),
    ("Issue date", "Date de delivrance"),
    ("Registration year", "Registration year"),
    ("Age", "Age"),
    ("Bond", "Bond"),
    ("Subcategory count", "Subcategory count"),
    ("Phone", "Numero de telephone"),
    ("Email", "Courriel"),
    ("Address", "Adresse"),
    ("Score", "Score"),
]


def to_output(df: pd.DataFrame) -> pd.DataFrame:
    out = pd.DataFrame({new: df[old] if old in df.columns else "" for new, old in OUTPUT_MAP})
    return out


def step8_write(shortlist: pd.DataFrame, corridor: pd.DataFrame, p: Params,
                counts: list, path_or_buffer) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font

    params_rows = [(k, ", ".join(v) if isinstance(v, list) else v)
                   for k, v in asdict(p).items()]
    params_rows.append(("", ""))
    params_rows += [(f"Row count: {label}", n) for label, n in counts]
    params_df = pd.DataFrame(params_rows, columns=["Parameter", "Value"])

    with pd.ExcelWriter(path_or_buffer, engine="openpyxl") as xw:
        to_output(shortlist).to_excel(xw, sheet_name="Targets", index=False)
        to_output(corridor).to_excel(xw, sheet_name="Corridor", index=False)
        params_df.to_excel(xw, sheet_name="Parameters", index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial")
            for i, col in enumerate(ws.columns, start=1):
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
            ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def run(rbq_source, nom_source, p: Params, out, log_fn=print):
    """Run all steps. Returns (shortlist, corridor, counts). Raises EmptyResult."""
    counts = []

    def log(label, n):
        counts.append((label, int(n)))
        log_fn(f"{label}: {n:,}")

    df = step1_read_rbq(rbq_source, p, log)
    df = step2_filter(df, p, log)
    corridor = step3_territory(df, p, log)
    del df
    if p.USE_REQ_JOIN and nom_source is not None:
        corridor = step4_join_req(corridor, nom_source, log)
    else:
        log("STEP 4 - REQ join skipped", len(corridor))
    corridor = step5_calculate(corridor, p, log)
    corridor = step6_score(corridor, p, log)
    shortlist = step7_rank(corridor, p, log)
    step8_write(shortlist, corridor, p, counts, out)
    return shortlist, corridor, counts
